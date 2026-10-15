"""Text extraction from PDF, DOCX, TXT and image documents.

Failures are explicit:

* ``DocumentReadError``: the file cannot be parsed (corrupt, encrypted, wrong
  content). The document is marked ``failed`` with the message.
* ``OCRUnavailableError``: the document needs OCR but no OCR engine is
  installed. The document is marked ``requires_review``.

The extractor never falls back to reading raw bytes as text.
"""

from __future__ import annotations

import logging
import os
from pathlib import Path
from typing import Any, Protocol

from app.processing.ocr import OCRBox, boxes_from_paddle, read_page, spans_from_boxes
from app.processing.schemas import ExtractedSpan, NormalizedDocument, NormalizedPage

logger = logging.getLogger(__name__)

_MIN_WORDS_PER_PAGE = 10
_MIN_WORDS_WITH_IMAGE = 80
OCR_MODEL_DIR = Path(__file__).resolve().parents[2] / "models" / "ocr"
_ENGLISH_MODELS = ("en_PP-OCRv3_det_infer.onnx", "en_PP-OCRv4_rec_infer.onnx", "en_dict.txt")


class DocumentReadError(RuntimeError):
    """The document could not be parsed. The message is safe to show to the user."""


class OCRUnavailableError(RuntimeError):
    """The document needs OCR but no OCR engine is available."""


class OCRAdapter(Protocol):
    def extract(self, image: Any, page_number: int) -> list[ExtractedSpan]: ...


class PaddleOCRAdapter:
    """Lazily loaded PaddleOCR 2.x engine (optional dependency) with layout rebuilding.

    ``preprocess="auto"`` applies only the corrections a page needs (upscaling,
    contrast, orientation, deskew; see ``app.processing.ocr``); ``"none"`` reads
    the image as given.
    """

    name = "paddle"

    def __init__(self, language: str = "en", preprocess: str = "auto") -> None:
        self.language = language
        self.preprocess = preprocess
        self._engine: Any | None = None
        self.last_steps: list[str] = []

    def _get_engine(self) -> Any:
        if self._engine is None:
            try:
                # Load pyclipper before Paddle's native libraries initialize zlib.
                # PaddlePaddle 2.6 can otherwise corrupt the later extension import.
                import pyclipper  # type: ignore[import-not-found]  # noqa: F401
                from paddleocr import PaddleOCR  # type: ignore[import-not-found]
            except ImportError as error:
                raise OCRUnavailableError(
                    "This document needs OCR, but the OCR engine is not installed "
                    "(see backend/requirements-ocr.txt)."
                ) from error
            self._engine = PaddleOCR(
                lang=self.language,
                use_angle_cls=False,
                show_log=False,
                cpu_threads=max(1, int(os.getenv("OCR_CPU_THREADS", "2"))),
                rec_batch_num=1,
                mkldnn_cache_capacity=2,
            )
        return self._engine

    def read(self, image: Any) -> list[OCRBox]:
        return boxes_from_paddle(self._get_engine().ocr(image, cls=False))

    def extract(self, image: Any, page_number: int) -> list[ExtractedSpan]:
        self._get_engine()  # raise OCRUnavailableError before any image work
        reading = read_page(self, image, self.preprocess)
        self.last_steps = reading.steps
        if reading.steps:
            logger.info("ocr_page_corrections", extra={"page": page_number, "steps": reading.steps})
        return spans_from_boxes(reading.boxes, page_number)


class RapidOCRAdapter(PaddleOCRAdapter):
    """The PP-OCR detection and recognition models run on ONNX Runtime (RapidOCR).

    Same models and the same layout/orientation/deskew pipeline as the Paddle
    adapter, but a fraction of the memory: PaddlePaddle's CPU inference keeps
    buffers for every distinct crop shape (measured 2.5 GB for one page),
    which got the worker killed under a 3.5 GB container limit.
    """

    name = "rapidocr"

    def _get_engine(self) -> Any:
        if self._engine is None:
            try:
                from rapidocr_onnxruntime import RapidOCR  # type: ignore[import-not-found]
            except ImportError as error:
                raise OCRUnavailableError(
                    "This document needs OCR, but the OCR engine is not installed "
                    "(see backend/requirements-ocr.txt)."
                ) from error
            threads = max(1, int(os.getenv("OCR_CPU_THREADS", "2")))
            options: dict[str, Any] = {
                "use_cls": False,
                "intra_op_num_threads": threads,
                "inter_op_num_threads": 1,
                # PaddleOCR's detection sizing (longest side 960)
                "det_limit_side_len": 960,
                "det_limit_type": "max",
            }
            model_dir = Path(os.getenv("OCR_MODEL_DIR") or OCR_MODEL_DIR)
            det, rec, keys = (model_dir / name for name in _ENGLISH_MODELS)
            if det.exists() and rec.exists() and keys.exists():
                # English PP-OCR models (converted from PaddleOCR's "en" models):
                # RapidOCR's bundled Chinese recogniser drops the spaces between
                # English words ("DateofResult"), which breaks label matching.
                options.update(
                    det_model_path=str(det), rec_model_path=str(rec), rec_keys_path=str(keys)
                )
            else:
                logger.warning("ocr_english_models_missing", extra={"model_dir": str(model_dir)})
            self._engine = RapidOCR(**options)
        return self._engine

    def read(self, image: Any) -> list[OCRBox]:
        result, _elapsed = self._get_engine()(image, use_cls=False)
        # RapidOCR returns [quad, text, score]; reuse the Paddle conversion
        return boxes_from_paddle([[[quad, (text, score)] for quad, text, score in result or []]])


class NoOCRAdapter:
    """OCR disabled by configuration (``OCR_ENGINE=none``)."""

    name = "none"
    last_steps: list[str] = []

    def extract(self, image: Any, page_number: int) -> list[ExtractedSpan]:
        raise OCRUnavailableError("This document needs OCR, but OCR is disabled (OCR_ENGINE=none).")


def build_ocr_adapter(
    engine: str = "rapidocr", language: str = "en", preprocess: str = "auto"
) -> Any:
    if engine == "none":
        return NoOCRAdapter()
    if engine == "paddle":
        return PaddleOCRAdapter(language=language, preprocess=preprocess)
    return RapidOCRAdapter(language=language, preprocess=preprocess)


class UniversalDocumentExtractor:
    """Dispatches on file extension: .pdf, .docx, .txt, .png, .jpg, .jpeg."""

    def __init__(self, ocr: OCRAdapter | None = None, resolution: int = 150) -> None:
        self.ocr = ocr or PaddleOCRAdapter()
        self.resolution = resolution

    def extract(self, path: str, document_name: str = "") -> NormalizedDocument:
        suffix = Path(document_name or path).suffix.lower() or Path(path).suffix.lower()
        if suffix == ".pdf":
            return self._extract_pdf(path)
        if suffix == ".docx":
            return self._extract_docx(path)
        if suffix == ".txt":
            return self._extract_txt(path)
        if suffix in (".png", ".jpg", ".jpeg"):
            return self._extract_image(path)
        raise DocumentReadError(f"Unsupported document type: {suffix or 'unknown'}")

    # ------------------------------------------------------------------ PDF

    def _extract_pdf(self, path: str) -> NormalizedDocument:
        try:
            import pdfplumber
        except ImportError as error:  # pragma: no cover - dependency is required
            raise RuntimeError("pdfplumber is not installed") from error
        pages: list[NormalizedPage] = []
        warnings: list[str] = []
        ocr_missing = False
        try:
            with pdfplumber.open(path) as pdf:
                if not pdf.pages:
                    raise DocumentReadError("The PDF has no pages.")
                for number, page in enumerate(pdf.pages, start=1):
                    spans = _direct_spans(page, number)
                    direct_words = _word_count(spans)
                    # A scanned card pasted under a typed header has some text but
                    # its content is in the image: OCR it and add what is new.
                    pasted_scan = (
                        direct_words < _MIN_WORDS_WITH_IMAGE and _image_coverage(page) >= 0.35
                    )
                    if direct_words < _MIN_WORDS_PER_PAGE or pasted_scan:
                        try:
                            image = page.to_image(resolution=self.resolution).original
                            ocr_spans = self.ocr.extract(image, number)
                        except OCRUnavailableError:
                            ocr_missing = True
                            ocr_spans = []
                        if direct_words < _MIN_WORDS_PER_PAGE:
                            if _word_count(ocr_spans) > direct_words:
                                spans = ocr_spans
                        else:
                            known = {_compact(span.text) for span in spans}
                            spans = spans + [
                                span for span in ocr_spans if _compact(span.text) not in known
                            ]
                    pages.append(
                        NormalizedPage(
                            page_number=number, width=page.width, height=page.height, spans=spans
                        )
                    )
        except DocumentReadError:
            raise
        except Exception as error:
            logger.info("pdf_parse_failed", extra={"error_type": type(error).__name__})
            raise DocumentReadError(
                "The PDF could not be read. It may be corrupted or password-protected."
            ) from error
        document = NormalizedDocument(pages=pages)
        if ocr_missing:
            if not document.text.strip():
                raise OCRUnavailableError(
                    "This PDF is scanned and needs OCR, but the OCR engine is not installed."
                )
            warnings.append(
                "Some pages look scanned and were skipped because OCR is not installed."
            )
        document.warnings = warnings
        return document

    # ------------------------------------------------------------------ DOCX

    def _extract_docx(self, path: str) -> NormalizedDocument:
        try:
            import docx  # python-docx
        except ImportError as error:  # pragma: no cover - dependency is required
            raise RuntimeError("python-docx is not installed") from error
        try:
            word = docx.Document(path)
        except Exception as error:
            raise DocumentReadError("The Word document could not be read.") from error
        spans = [
            ExtractedSpan(page_number=1, text=paragraph.text.strip(), source_method="direct-parse")
            for paragraph in word.paragraphs
            if paragraph.text.strip()
        ]
        for table in word.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                line = " | ".join(cell for cell in cells if cell)
                if line:
                    spans.append(
                        ExtractedSpan(
                            page_number=1, text=line, source_method="direct-parse", is_table=True
                        )
                    )
        return NormalizedDocument(pages=[NormalizedPage(page_number=1, spans=spans)])

    # ------------------------------------------------------------------ TXT

    def _extract_txt(self, path: str) -> NormalizedDocument:
        try:
            content = Path(path).read_bytes().decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise DocumentReadError("The text file is not valid UTF-8.") from error
        spans = [
            ExtractedSpan(page_number=1, text=line.strip(), source_method="direct-parse")
            for line in content.splitlines()
            if line.strip()
        ]
        return NormalizedDocument(pages=[NormalizedPage(page_number=1, spans=spans)])

    # ------------------------------------------------------------------ images

    def _extract_image(self, path: str) -> NormalizedDocument:
        from PIL import Image, UnidentifiedImageError

        try:
            with Image.open(path) as image:
                image.load()
                width, height = image.size
                spans = self.ocr.extract(image.convert("RGB"), page_number=1)
        except UnidentifiedImageError as error:
            raise DocumentReadError("The image could not be read.") from error
        return NormalizedDocument(
            pages=[
                NormalizedPage(page_number=1, width=float(width), height=float(height), spans=spans)
            ]
        )


def _direct_spans(page: Any, page_number: int) -> list[ExtractedSpan]:
    spans = [
        ExtractedSpan(page_number=page_number, text=line.strip(), source_method="direct-parse")
        for line in (page.extract_text() or "").splitlines()
        if line.strip()
    ]
    for table in page.extract_tables() or []:
        rows = [" | ".join(str(cell or "").strip() for cell in row) for row in table]
        text = "\n".join(row for row in rows if row.strip(" |"))
        if text:
            spans.append(
                ExtractedSpan(
                    page_number=page_number, text=text, source_method="direct-parse", is_table=True
                )
            )
    return spans


def _compact(text: str) -> str:
    return "".join(ch for ch in text.casefold() if ch.isalnum())


def _image_coverage(page: Any) -> float:
    """Share of the page area covered by embedded images (0..1)."""
    try:
        area = float(page.width) * float(page.height)
        covered = sum(
            max(0.0, float(i["x1"]) - float(i["x0"]))
            * max(0.0, float(i["bottom"]) - float(i["top"]))
            for i in page.images or []
        )
    except (KeyError, TypeError, ValueError):
        return 0.0
    return min(1.0, covered / area) if area else 0.0


def _word_count(spans: list[ExtractedSpan]) -> int:
    return sum(1 for span in spans for word in span.text.split() if any(c.isalnum() for c in word))
