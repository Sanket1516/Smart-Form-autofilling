# -*- coding: utf-8 -*-
"""Small python-docx toolkit that implements the college report rules.

College rules (Report Writing Instructions.pdf):
  Times New Roman, body 12 pt, chapter titles 24 pt, line spacing 1.5,
  margins L 1.5" R 1.0" T 1.0" B 1.0", main headings UPPER CASE 12 pt,
  sub-headings Title Case 12 pt, roman page numbers for front matter,
  decimal from Chapter 1, every chapter has an introduction and a summary.
"""
import os
import re

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
FONT = "Times New Roman"
PAGE_W, PAGE_H = 8.27, 11.69
ML, MR, MT, MB = 1.5, 1.0, 1.0, 1.0
TEXT_W = PAGE_W - ML - MR  # 5.77
NAVY = RGBColor(0x1F, 0x3A, 0x66)


def _set_font(style_or_run, name=FONT, size=None, bold=None, italic=None, color=None):
    f = style_or_run.font
    f.name = name
    rpr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run.element, "get_or_add_rPr") else style_or_run.element.rPr
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color


def shade(paragraph_or_cell, fill):
    el = paragraph_or_cell._p.get_or_add_pPr() if hasattr(paragraph_or_cell, "_p") else paragraph_or_cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el.append(shd)


def add_field(paragraph, instr, placeholder="1", bold=None):
    def fc(t):
        r = paragraph.add_run()
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), t)
        r._r.append(e)
        return r
    fc("begin")
    r = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)
    fc("separate")
    rr = paragraph.add_run(placeholder)
    if bold is not None:
        rr.bold = bold
    fc("end")


class Report:
    def __init__(self, figs, tables):
        self.doc = Document()
        self.figs = figs      # key -> dict(file, caption, chapter, num)
        self.tabs = tables    # key -> dict(caption, chapter, num)
        self.fig_used, self.tab_used, self.fig_ref, self.tab_ref = set(), set(), set(), set()
        self.chapter_no = 0
        self._setup()

    # ------------------------------------------------------------ setup
    def _setup(self):
        d = self.doc
        s = d.sections[0]
        self._page(s, portrait=True)
        st = d.styles
        n = st["Normal"]
        _set_font(n, size=12)
        pf = n.paragraph_format
        pf.line_spacing = 1.5
        pf.space_after = Pt(6)
        pf.space_before = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.widow_control = True

        h1 = st["Heading 1"]
        _set_font(h1, size=24, bold=True, color=RGBColor(0, 0, 0))
        h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h1.paragraph_format.space_before = Pt(60)
        h1.paragraph_format.space_after = Pt(24)
        h1.paragraph_format.line_spacing = 1.15
        h1.paragraph_format.keep_with_next = True
        h1.paragraph_format.page_break_before = True
        h2 = st["Heading 2"]
        _set_font(h2, size=12, bold=True, color=RGBColor(0, 0, 0))
        h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h2.paragraph_format.space_before = Pt(14)
        h2.paragraph_format.space_after = Pt(6)
        h2.paragraph_format.keep_with_next = True
        h2.paragraph_format.line_spacing = 1.5
        h3 = st["Heading 3"]
        _set_font(h3, size=12, bold=True, color=RGBColor(0, 0, 0))
        h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h3.paragraph_format.space_before = Pt(8)
        h3.paragraph_format.space_after = Pt(4)
        h3.paragraph_format.keep_with_next = True
        h3.paragraph_format.line_spacing = 1.5
        h4 = st["Heading 4"]
        _set_font(h4, size=12, bold=True, italic=True, color=RGBColor(0, 0, 0))
        h4.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h4.paragraph_format.space_before = Pt(6)
        h4.paragraph_format.space_after = Pt(3)
        h4.paragraph_format.keep_with_next = True
        h4.paragraph_format.line_spacing = 1.5

        def pstyle(name, size=12, bold=False, italic=False, align=None, ls=1.5, before=0, after=6, base="Normal", keep_next=False):
            s_ = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            s_.base_style = st[base]
            _set_font(s_, size=size, bold=bold, italic=italic, color=RGBColor(0, 0, 0))
            s_.paragraph_format.line_spacing = ls
            s_.paragraph_format.space_before = Pt(before)
            s_.paragraph_format.space_after = Pt(after)
            if align is not None:
                s_.paragraph_format.alignment = align
            s_.paragraph_format.keep_with_next = keep_next
            return s_

        fh = pstyle("FrontHeading", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, ls=1.15, before=0, after=18, keep_next=True)
        fh.paragraph_format.page_break_before = True
        pstyle("FigureCaption", size=12, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, ls=1.15, before=4, after=14)
        pstyle("TableCaption", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, ls=1.15, before=10, after=4, keep_next=True)
        pstyle("TableText", size=10.5, ls=1.0, after=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        pstyle("Centered", align=WD_ALIGN_PARAGRAPH.CENTER, ls=1.5)
        pstyle("IntroPara", italic=False, ls=1.5)
        c = pstyle("CodeBlock", size=8, ls=1.0, after=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        _set_font(c, name="Courier New", size=8, color=RGBColor(0, 0, 0))
        for name, size, bold, indent in (("toc 1", 12, True, 0), ("toc 2", 12, False, 0.3), ("toc 3", 11, False, 0.6)):
            try:
                t = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            except Exception:
                t = st[name]
            t.base_style = st["Normal"]
            _set_font(t, size=size, bold=bold, color=RGBColor(0, 0, 0))
            t.paragraph_format.line_spacing = 1.15
            t.paragraph_format.space_after = Pt(3)
            t.paragraph_format.left_indent = Inches(indent)
            t.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        lb = st["List Bullet"]
        _set_font(lb, size=12)
        lb.paragraph_format.line_spacing = 1.5
        lb.paragraph_format.space_after = Pt(3)
        lb.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        ln = st["List Number"]
        _set_font(ln, size=12)
        ln.paragraph_format.line_spacing = 1.5
        ln.paragraph_format.space_after = Pt(3)
        ln.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        # footer page number
        self._footer(s)

    def _page(self, s, portrait=True):
        if portrait:
            s.orientation = WD_ORIENT.PORTRAIT
            s.page_width, s.page_height = Inches(PAGE_W), Inches(PAGE_H)
            s.left_margin, s.right_margin = Inches(ML), Inches(MR)
        else:
            s.orientation = WD_ORIENT.LANDSCAPE
            s.page_width, s.page_height = Inches(PAGE_H), Inches(PAGE_W)
            s.left_margin, s.right_margin = Inches(1.0), Inches(1.0)
        s.top_margin, s.bottom_margin = Inches(MT), Inches(MB)
        s.footer_distance = Inches(0.45)

    def _footer(self, s):
        s.footer.is_linked_to_previous = False
        p = s.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        add_field(p, "PAGE", "1")
        for r in p.runs:
            _set_font(r, size=11)

    def _pgnum(self, s, fmt=None, start=None):
        sect = s._sectPr
        for e in sect.findall(qn("w:pgNumType")):
            sect.remove(e)
        pg = OxmlElement("w:pgNumType")
        if fmt:
            pg.set(qn("w:fmt"), fmt)
        if start is not None:
            pg.set(qn("w:start"), str(start))
        sect.append(pg)

    def _no_footer(self, s):
        s.footer.is_linked_to_previous = False
        for p in s.footer.paragraphs:
            for r in list(p.runs):
                r._r.getparent().remove(r._r)

    def new_section(self, fmt=None, start=None, portrait=True, footer=True):
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._page(s, portrait)
        self._pgnum(s, fmt, start)
        if footer:
            self._footer(s)
        else:
            self._no_footer(s)
        return s

    # ------------------------------------------------------------ text
    def p(self, text="", style=None, align=None, bold=False, italic=False, keep_next=False, size=None, after=None):
        para = self.doc.add_paragraph(style=style)
        if text:
            self._runs(para, text, bold, italic, size)
        if align is not None:
            para.alignment = align
        if keep_next:
            para.paragraph_format.keep_with_next = True
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        return para

    def _runs(self, para, text, bold=False, italic=False, size=None):
        """Supports **bold** and *italic* and `code` inline markup."""
        parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
        for part in parts:
            if not part:
                continue
            if part.startswith("**"):
                r = para.add_run(part[2:-2])
                r.bold = True
                r.italic = italic or None
            elif part.startswith("`"):
                r = para.add_run(part[1:-1])
                _set_font(r, name="Courier New", size=10.5)
                r.bold = bold or None
            elif part.startswith("*") and len(part) > 2:
                r = para.add_run(part[1:-1])
                r.italic = True
                r.bold = bold or None
            else:
                r = para.add_run(part)
                r.bold = bold or None
                r.italic = italic or None
            if size and not part.startswith("`"):
                r.font.size = Pt(size)

    def chapter(self, num, title):
        self.chapter_no = num
        para = self.doc.add_paragraph(style="Heading 1")
        para.add_run(f"Chapter {num}")
        para.add_run().add_break(WD_BREAK.LINE)
        para.add_run(title)
        return para

    def h1_plain(self, title):
        return self.doc.add_paragraph(title, style="Heading 1")

    def front(self, title):
        return self.doc.add_paragraph(title, style="FrontHeading")

    def h2(self, text):
        return self.doc.add_paragraph(text.upper(), style="Heading 2")

    def h3(self, text):
        return self.doc.add_paragraph(text, style="Heading 3")

    def h4(self, text):
        return self.doc.add_paragraph(text, style="Heading 4")

    def bullets(self, items, lead_bold=True):
        for it in items:
            para = self.doc.add_paragraph(style="List Bullet")
            if lead_bold and "::" in it:
                a, b = it.split("::", 1)
                r = para.add_run(a.strip() + ": ")
                r.bold = True
                self._runs(para, b.strip())
            else:
                self._runs(para, it)

    def numbered(self, items):
        # manual numbering: avoids list-numbering continuation problems across the document
        for i, it in enumerate(items, 1):
            para = self.doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.35)
            para.paragraph_format.first_line_indent = Inches(-0.3)
            para.paragraph_format.space_after = Pt(3)
            self._runs(para, f"{i}.  " + it)

    def intro(self, text):
        para = self.p(text, style="IntroPara")
        return para

    def summary(self, text):
        self.h2("Chapter Summary")
        return self.p(text)

    # ------------------------------------------------------------ refs
    def f(self, key):
        self.fig_ref.add(key)
        return f"Figure {self.figs[key]['chapter']}.{self.figs[key]['num']}"

    def t(self, key):
        self.tab_ref.add(key)
        return f"Table {self.tabs[key]['chapter']}.{self.tabs[key]['num']}"

    # ------------------------------------------------------------ figures
    def fig(self, key, max_w=None, force_landscape=None, max_h=None):
        meta = self.figs[key]
        self.fig_used.add(key)
        path = os.path.join(HERE, "figures", meta["file"])
        w_px, h_px = Image.open(path).size
        aspect = w_px / h_px
        cap = f"Figure {meta['chapter']}.{meta['num']}: {meta['caption']}"
        # choose orientation that maximises scale
        dpi = meta.get("dpi", 200)
        nat_w = w_px / dpi
        port_w = min(TEXT_W, 6.4 * aspect, nat_w if meta.get("no_upscale", True) else 99)
        land_w = min(9.19, 5.0 * aspect, nat_w if meta.get("no_upscale", True) else 99)
        land = force_landscape if force_landscape is not None else (land_w > port_w * 1.25)
        if land:
            self.new_section(portrait=False)
            width = land_w
        else:
            width = port_w
        if max_w:
            width = min(width, max_w)
        if max_h:
            width = min(width, max_h * aspect)
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.paragraph_format.line_spacing = 1.0
        para.paragraph_format.space_before = Pt(6)
        para.add_run().add_picture(path, width=Inches(width))
        cp = self.doc.add_paragraph(cap, style="FigureCaption")
        # bold label
        txt = cp.runs[0].text
        label, rest = txt.split(":", 1)
        cp.runs[0].text = label + ":"
        cp.runs[0].bold = True
        cp.add_run(rest)
        if land:
            self.new_section(portrait=True)

    # ------------------------------------------------------------ tables
    def table(self, key, rows, widths, font=10.5, header=True, align=None, bold_first_col=False):
        meta = self.tabs[key]
        self.tab_used.add(key)
        cap = self.doc.add_paragraph(style="TableCaption")
        r = cap.add_run(f"Table {meta['chapter']}.{meta['num']}: ")
        r.bold = True
        cap.add_run(meta["caption"])
        tbl = self.doc.add_table(rows=len(rows), cols=len(widths))
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tblPr = tbl._tbl.tblPr
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblPr.append(lay)
        for ri, row in enumerate(rows):
            tr = tbl.rows[ri]
            trPr = tr._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            trPr.append(cs)
            if header and ri == 0:
                th = OxmlElement("w:tblHeader")
                trPr.append(th)
            for ci, val in enumerate(row):
                cell = tr.cells[ci]
                cell.width = Inches(widths[ci])
                cell.paragraphs[0].style = "TableText"
                para = cell.paragraphs[0]
                para.paragraph_format.space_before = Pt(2)
                para.paragraph_format.space_after = Pt(2)
                if align and align[ci] == "c":
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                lines = str(val).split("\n")
                for li, line in enumerate(lines):
                    if li:
                        para = cell.add_paragraph(style="TableText")
                        para.paragraph_format.space_after = Pt(2)
                    self._runs(para, line, bold=(header and ri == 0) or (bold_first_col and ci == 0 and ri > 0), size=font)
                if header and ri == 0:
                    shade(cell, "D9E2F3")
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(6)
        spacer.paragraph_format.line_spacing = 0.6
        return tbl

    def code(self, text, caption=None):
        if caption:
            self.p(caption, bold=True, keep_next=True, after=2)
        lines = text.rstrip("\n").split("\n")
        for i, ln in enumerate(lines):
            para = self.doc.add_paragraph(style="CodeBlock")
            para.paragraph_format.keep_together = True
            shade(para, "F2F2F2")
            para.paragraph_format.left_indent = Inches(0.05)
            run = para.add_run(ln if ln else " ")
            _set_font(run, name="Courier New", size=8)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def toc(self, instr):
        para = self.doc.add_paragraph()
        para.paragraph_format.line_spacing = 1.15
        add_field(para, instr, "Right-click and choose Update Field to build this list.")
        return para

    def save(self, path):
        self.doc.save(path)
