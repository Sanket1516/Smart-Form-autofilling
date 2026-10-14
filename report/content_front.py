# -*- coding: utf-8 -*-
import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from docbuild import HERE, TEXT_W

TITLE = "Agentic AI Framework for Context-Aware Semantic Understanding and Intelligent Assistance of Dynamic Web Forms"
SHORT = "Smart Form Auto-Filling System"
TEAM = [("Ashmit Patil", "E246", "70552300048"), ("Sanket Mali", "E265", "70552300067"), ("Pinak Salunke", "E247", "70552300049")]
GUIDE = "Dr. Nitin Choubey"
GUIDE_DESIG = "Professor, Department of Computer Science"
HOD = "Dr. Nitin S. Choubey"
DEAN = "Dr. Venkatadri Marriboyina"
SESSION = "2026-27"
C = WD_ALIGN_PARAGRAPH.CENTER


def line(R, text, size=12, bold=False, after=6, italic=False, before=0):
    p = R.p("", align=C, after=after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(before)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    return p


def sign_row(R, left, right, left_sub="", right_sub=""):
    tbl = R.doc.add_table(rows=1, cols=2)
    tbl.autofit = False
    for i, (a, b) in enumerate(((left, left_sub), (right, right_sub))):
        cell = tbl.rows[0].cells[i]
        cell.width = Inches(TEXT_W / 2)
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.15
        p.alignment = C
        p.add_run("______________________\n").font.size = Pt(12)
        r = p.add_run(a)
        r.bold = True
        if b:
            p.add_run("\n" + b)
    R.p("", after=4)


def cover(R):
    d = R.doc
    line(R, TITLE, 24, True, after=8, before=4)
    line(R, f"({SHORT})", 14, False, after=14, italic=True)
    line(R, "Submitted in partial fulfillment of the requirement for the award of", 12, False, after=0)
    line(R, "Degree of Bachelor of Technology in", 14, True, after=0)
    line(R, "Computer Science", 14, True, after=10)
    line(R, "Submitted To", 12, True, after=4)
    logo = os.path.join(HERE, "figures", "nmims_logo.png")
    p = R.p("", align=C, after=4)
    p.add_run().add_picture(logo, width=Inches(1.35))
    line(R, "SVKM’s NMIMS,", 12, True, after=0)
    line(R, "Mukesh Patel School of Technology Management & Engineering,", 12, True, after=0)
    line(R, "Shirpur Campus (M.H.)", 12, True, after=10)
    line(R, "Submitted By:", 12, True, after=4)
    for name, roll, enr in TEAM:
        line(R, f"{name} ({roll})   Enrollment No – {enr}", 14, True, after=2)
    line(R, "Under The Supervision Of:", 12, True, after=4, before=8)
    line(R, GUIDE, 14, True, after=0)
    line(R, f"({GUIDE_DESIG})", 14, False, after=10)
    line(R, "DEPARTMENT OF COMPUTER SCIENCE", 14, True, after=0)
    line(R, "Mukesh Patel School of Technology Management & Engineering", 12, True, after=0)
    line(R, f"ACADEMIC SESSION: {SESSION}", 12, True, after=0)


def certificate(R):
    R.front("CERTIFICATE")
    p = R.p("")
    p.add_run("This is to certify that the project entitled ")
    p.add_run(f"“{TITLE}”").bold = True
    p.add_run(" has been done by")
    for name, roll, enr in TEAM:
        q = R.p("", align=C, after=0)
        q.add_run(f"“{name}” (Roll No.: {roll[1:]}; Enrollment No.: {enr})").bold = True
    R.p("under my guidance and supervision and has been submitted in partial fulfillment of the degree of "
        "“Bachelor of Technology in Computer Science” of SVKM’s NMIMS (Deemed-to-be University), MPSTME Shirpur Campus (M.H.), India.")
    R.p("", after=14)
    sign_row(R, "Project Mentor", "Examiner", f"({GUIDE}, Internal Guide)", "")
    R.p("Date:", after=0)
    R.p("Place: Shirpur", after=6)
    sign_row(R, "H.O.D.", "", f"({HOD})", "")
    line(R, "DEPARTMENT OF COMPUTER SCIENCE", 14, True, after=0, before=10)
    line(R, "Mukesh Patel School of Technology Management & Engineering", 12, True, after=0)


def approval(R):
    R.front("CERTIFICATE OF APPROVAL")
    R.p(f"The project entitled “{TITLE}” submitted by Ashmit Patil (70552300048), Sanket Mali (70552300067) and "
        "Pinak Salunke (70552300049) of B.Tech (Computer Science), Academic Session 2026-27, MPSTME Shirpur Campus, SVKM’s NMIMS, has been "
        "examined and is approved for the award of the degree of Bachelor of Technology in Computer Science.")
    R.p("The examiners record their approval below after the project presentation and viva-voce examination held on "
        "____________________.", after=24)
    sign_row(R, "Internal Examiner", "External Examiner", "Name: ______________________", "Name: ______________________")
    R.p("Date: ______________", after=0)
    R.p("Place: Shirpur", after=18)
    sign_row(R, "Project Guide", "Head of Department", f"({GUIDE})", f"({HOD})")


def declaration(R):
    R.front("DECLARATION")
    R.p(f"We, the undersigned students of B.Tech (Computer Science), Academic Session {SESSION}, MPSTME Shirpur Campus, SVKM’s NMIMS, declare "
        f"that the project report entitled “{TITLE}” is a record of the work carried out by us under the guidance of {GUIDE}. "
        "The software, diagrams, test results and screenshots in this report are produced from the implementation developed by the project team.")
    R.p("To the best of our knowledge, this report has not been submitted, in part or in full, to any other university or institution for the award of "
        "any degree or diploma. Information obtained from published or unpublished sources has been acknowledged in the text and listed in the References.")
    R.p("", after=24)
    sign_row(R, "Ashmit Patil", "Sanket Mali", "Enrollment No. 70552300048", "Enrollment No. 70552300067")
    sign_row(R, "Pinak Salunke", "", "Enrollment No. 70552300049", "")
    R.p("Date: ______________", after=0)
    R.p("Place: Shirpur")


def acknowledgement(R):
    R.front("ACKNOWLEDGEMENT")
    R.p("After the completion of this capstone project work, words are not enough to express our feelings about all those who helped us to reach "
        "our goal; feeling above this is our indebtedness to The Almighty for providing us this moment in life.")
    R.p(f"It is a great pleasure and a moment of immense satisfaction for us to express our profound gratitude to {GUIDE}, Professor, Computer Science "
        "Department, MPSTME, Shirpur, whose constant encouragement enabled us to work enthusiastically. Their perpetual motivation, patience and "
        "expertise in discussion during the progress of the project work have benefited us to an extent which is beyond expression. Their depth and "
        "breadth of knowledge of the Computer Science field made us realise that theoretical knowledge always helps to develop efficient operational "
        "software, which is a blend of all core subjects of the field. We are highly indebted to them for their invaluable guidance and ever-ready "
        "support in the successful completion of this project in time.")
    R.p(f"We express our sincere thanks and gratitude to {HOD}, Head of Department, Computer Science, MPSTME, Shirpur, for providing the necessary "
        "infrastructure and help to complete the project work successfully.")
    R.p(f"We also extend our deepest gratitude to {DEAN}, Associate Dean, SVKM’s NMIMS, Shirpur Campus, for providing all the necessary facilities and "
        "an encouraging environment to bring out the best of our endeavours.")
    R.p("We sincerely wish to express our grateful thanks to all members of the staff of the Computer Science department and to all those who have "
        "imparted technical knowledge of computer technology to us during the various stages of B.Tech. Computer Science. We also thank our friends, "
        "who contributed directly or indirectly to this capstone project work, and the open-source communities whose tools made it possible.")
    R.p("", after=18)
    for name, roll, enr in TEAM:
        p = R.p(name, align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
        p.paragraph_format.line_spacing = 1.15


def abstract(R):
    R.front("ABSTRACT")
    R.p("Filling web forms is a repetitive task in which people re-type the same identity, education, address and employment details on portals for "
        "admissions, banking, government services and recruitment. Native browser autofill only replays values typed earlier, cannot read a user’s "
        "documents, records no evidence for a value, and does not warn when two documents disagree. This report presents the design, implementation and "
        "verification of the Smart Form Auto-Filling System, a privacy-first assistant that builds a profile from the user’s own documents and uses it to "
        "recommend values for a web form, which the user reviews before anything is filled.")
    R.p("The system has three parts: a FastAPI backend with MongoDB and ChromaDB, a React dashboard, and a Chrome Manifest V3 extension. Uploaded PDF, DOCX, "
        "TXT and image files are parsed; facts are extracted by deterministic rules over a 91-concept YAML ontology and, optionally, by a large language "
        "model, then merged with provenance (origin, evidence text, page, confidence). Disagreements between documents are computed as conflicts for the "
        "user to resolve. A seven-node LangGraph pipeline classifies each form field, resolves a value from the profile or from retrieved document evidence, "
        "and applies confidence and sensitivity rules. Prompt-injection defences, tenant isolation by user identifier, token revocation, redacted logging and "
        "an on-demand extension permission model protect the user’s data. The extension never reads field values, never fills password, payment or "
        "one-time-code fields and never submits a form.")
    R.p("The implementation is verified by 196 automated tests, all of which passed in the recorded run: 146 backend tests with a mock LLM, 3 integration "
        "tests against a real ChromaDB service, 22 dashboard tests, 24 extension tests and one Playwright browser test on an isolated Docker Compose stack. "
        "Real-provider LLM tests were not run. The main contribution is an integrated, evidence-preserving and human-in-the-loop pipeline from document to "
        "filled form that is deployable with a single Compose file.")
    p = R.p("")
    p.add_run("Keywords: ").bold = True
    p.add_run("Web form autofill, document information extraction, retrieval-augmented generation, LangGraph, field ontology, prompt injection, "
              "Chrome Manifest V3, privacy.")


ABBREVIATIONS = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"), ("bcrypt", "Password hashing function based on the Blowfish cipher"),
    ("CI", "Continuous Integration"), ("CORS", "Cross-Origin Resource Sharing"), ("CSP", "Content Security Policy"), ("DFD", "Data Flow Diagram"),
    ("DOM", "Document Object Model"), ("E2E", "End to End"), ("ER", "Entity Relationship"), ("HTTP", "Hypertext Transfer Protocol"),
    ("JSON", "JavaScript Object Notation"), ("JWT", "JSON Web Token"), ("KYC", "Know Your Customer"), ("LLM", "Large Language Model"),
    ("MV3", "Manifest Version 3 (Chrome extension platform)"), ("NER", "Named Entity Recognition"), ("OCR", "Optical Character Recognition"),
    ("OTP", "One-Time Password"), ("PII", "Personally Identifiable Information"), ("RAG", "Retrieval-Augmented Generation"),
    ("REST", "Representational State Transfer"), ("SPA", "Single-Page Application"), ("TTL", "Time To Live"), ("UI", "User Interface"),
    ("UML", "Unified Modelling Language"), ("XSS", "Cross-Site Scripting"), ("YAML", "YAML Ain’t Markup Language"),
]
