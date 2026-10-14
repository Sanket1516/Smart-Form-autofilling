# -*- coding: utf-8 -*-
"""Assembles the B.Tech black book. Run: python build_report.py  (then finalize.ps1 updates fields and exports PDF)."""
import sys
sys.path.insert(0, ".")
import registry
from docbuild import Report
import content_front as F
import content_refs
import ch1, ch2, ch3, ch4a, ch4b, ch5, ch6, ch7, appendices
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = "Smart_Form_Black_Book_2026-27.docx"

figs, tabs = registry.build()
R = Report(figs, tabs)

# ---- cover (own section, no page number)
R._no_footer(R.doc.sections[0])
R._pgnum(R.doc.sections[0], "lowerRoman", 1)
F.cover(R)

# ---- front matter: roman numerals, certificate is page ii
R.new_section(fmt="lowerRoman", start=2)
F.certificate(R)
F.approval(R)
F.declaration(R)
F.acknowledgement(R)
F.abstract(R)
R.front("TABLE OF CONTENTS")
R.toc(r'TOC \o "1-3" \h \z \u')
R.front("LIST OF FIGURES")
R.toc(r'TOC \h \z \t "FigureCaption,1"')
R.front("LIST OF TABLES")
R.toc(r'TOC \h \z \t "TableCaption,1"')
R.front("LIST OF ABBREVIATIONS")
R.table("abbr", [["Abbreviation", "Expansion"]] + [list(x) for x in F.ABBREVIATIONS], [1.3, 4.45], font=10.5) if False else None
t = R.doc.add_table(rows=len(F.ABBREVIATIONS) + 1, cols=2)
t.style = "Table Grid"
for i, row in enumerate([("Abbreviation", "Expansion")] + F.ABBREVIATIONS):
    for j, v in enumerate(row):
        cell = t.rows[i].cells[j]
        cell.paragraphs[0].style = "TableText"
        r = cell.paragraphs[0].add_run(v)
        r.bold = (i == 0)

# ---- main matter: decimal numerals restart at 1
R.new_section(fmt="decimal", start=1)
ch1.build(R); ch2.build(R); ch3.build(R); ch4a.build(R); ch4b.build(R); ch5.build(R); ch6.build(R); ch7.build(R)
appendices.references(R)
appendices.build(R)

R.save(OUT)

# ---- consistency checks
unused_f = set(figs) - R.fig_used
unused_t = set(tabs) - R.tab_used
unref_f = R.fig_used - R.fig_ref
unref_t = R.tab_used - R.tab_ref
uncited = {k for k, _ in content_refs.REFS} - content_refs.CITED
print("figures", len(figs), "tables", len(tabs))
print("figures never inserted:", sorted(unused_f))
print("tables never inserted:", sorted(unused_t))
print("figures not referenced in text:", sorted(unref_f))
print("tables not referenced in text:", sorted(unref_t))
print("references never cited:", sorted(uncited))
