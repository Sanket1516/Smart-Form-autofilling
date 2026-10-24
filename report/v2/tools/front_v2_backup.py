# -*- coding: utf-8 -*-
"""Preliminary pages, laid out after 'Front Pages of Project Report (Black Book).doc'."""
import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from docbuild2 import HERE, TEXT_W

TITLE = "Agentic AI Framework for Context-Aware Semantic Understanding and Intelligent Assistance of Dynamic Web Forms"
SHORT = "Smart Form Auto-Filling System"
# (name, roll number, enrollment number) as recorded in the earlier submission
TEAM = [("Ashmit Patil", "246", "70552300048"), ("Sanket Mali", "265", "70552300067"), ("Pinak Salunke", "247", "70552300049")]
# Reviewer note: the guide's designation is Head of Department, not Professor.
GUIDE = "Dr. Nitin S. Choubey"
GUIDE_DESIG = "Head of Department, Computer Science"
HOD = "Dr. Nitin S. Choubey"
DEAN = "Dr. Venkatadri Marriboyina"
SESSION = "2026-27"
LOGO = os.path.join(HERE, "..", "figures", "nmims_logo.png")
C = WD_ALIGN_PARAGRAPH.CENTER


def line(R, text, size=12, bold=False, after=6, italic=False, before=0, underline=False):
    p = R.p("", align=C, after=after, before=before, ls=1.15)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.underline = underline
    return p


def sign_row(R, left, right, left_sub="", right_sub="", after=4):
    tbl = R.doc.add_table(rows=1, cols=2)
    tbl.autofit = False
    for i, (a, b) in enumerate(((left, left_sub), (right, right_sub))):
        cell = tbl.rows[0].cells[i]
        cell.width = Inches(TEXT_W / 2)
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.RIGHT
        if not a:
            continue
        p.add_run("____________________\n")
        r = p.add_run(a)
        r.bold = True
        if b:
            p.add_run("\n" + b)
    R.p("", after=after, ls=1.0)


def cover(R):
    line(R, TITLE, 24, True, after=6, before=0)
    line(R, f"({SHORT})", 14, False, after=14, italic=True)
    line(R, "Submitted in partial fulfillment of the requirement for the award of", 14, True, after=0)
    line(R, "Degree of Bachelor of Technology in", 14, True, after=0)
    line(R, "Computer Science", 14, True, after=10)
    line(R, "Submitted To", 12, True, after=4)
    p = R.p("", align=C, after=4, ls=1.0)
    p.add_run().add_picture(LOGO, width=Inches(1.3))
    line(R, "SVKM’s NMIMS,", 12, True, after=0)
    line(R, "Mukesh Patel School of Technology Management & Engineering,", 12, True, after=0)
    line(R, "Shirpur Campus (M.H.)", 12, True, after=10)
    line(R, "Submitted By:", 12, True, after=4)
    for name, roll, enr in TEAM:
        line(R, f"{name}   Enrollment No – {enr}", 14, True, after=1)
    line(R, "Under The Supervision Of:", 12, True, after=4, before=10)
    line(R, GUIDE, 14, True, after=0)
    line(R, f"({GUIDE_DESIG})", 14, False, after=18)
    line(R, "DEPARTMENT OF COMPUTER SCIENCE", 14, True, after=0)
    line(R, "Mukesh Patel School of Technology Management & Engineering", 12, True, after=0)
    line(R, f"ACADEMIC SESSION: {SESSION}", 12, True, after=0)


def certificate(R):
    h = R.front("CERTIFICATE")
    h.runs[0].underline = True
    p = R.p("", align=C, after=10, ls=1.0)
    p.add_run().add_picture(LOGO, width=Inches(1.1))
    p = R.p("")
    p.add_run("This is to certify that the project entitled ")
    p.add_run(f"“{TITLE}”").bold = True
    p.add_run(" has been done by")
    for name, roll, enr in TEAM:
        q = R.p("", align=C, after=0)
        r = q.add_run(f"“{name}” (Roll No.: {roll}; Enrollment No.: {enr})")
        r.bold = True
        r.font.size = Pt(14)
    R.p("under my guidance and supervision and has been submitted in partial fulfillment of the degree of "
        "“**Bachelor of Technology in Computer Science**” of SVKM’s NMIMS (Deemed-to-be University), MPSTME Shirpur Campus (M.H.), India.",
        before=6)
    R.p("", after=22, ls=1.0)
    sign_row(R, "Project Mentor", "Examiner", f"(Internal Guide)\n{GUIDE}", "", after=8)
    R.p("Date:", after=0)
    R.p("Place: Shirpur", after=0)
    sign_row(R, "", "H.O.D.", "", HOD, after=10)
    line(R, "DEPARTMENT OF COMPUTER SCIENCE", 14, True, after=0, before=6)
    line(R, "Mukesh Patel School of Technology Management & Engineering", 12, True, after=0)


def approval(R):
    h = R.front("CERTIFICATE OF APPROVAL")
    h.runs[0].underline = True
    names = ", ".join(f"{n} ({e})" for n, _, e in TEAM[:-1]) + f" and {TEAM[-1][0]} ({TEAM[-1][2]})"
    R.p(f"The project entitled “**{TITLE}**”, submitted by {names} of B.Tech. (Computer Science), Academic Session {SESSION}, "
        "MPSTME Shirpur Campus, SVKM’s NMIMS, has been examined and is approved in partial fulfillment of the requirements for the award of the degree "
        "of Bachelor of Technology in Computer Science.")
    R.p("The examiners record their approval below after the project presentation and viva-voce examination held on ____________________.", after=30)
    sign_row(R, "Internal Examiner", "External Examiner", "Name: ____________________", "Name: ____________________", after=24)
    sign_row(R, "Project Guide", "Head of Department", GUIDE, HOD, after=18)
    R.p("Date: ____________________", after=0)
    R.p("Place: Shirpur")


def declaration(R):
    h = R.front("DECLARATION")
    h.runs[0].underline = True
    R.p(f"We, the undersigned students of B.Tech. (Computer Science), Academic Session {SESSION}, MPSTME Shirpur Campus, SVKM’s NMIMS, declare "
        f"that the project report entitled “**{TITLE}**” is a record of the work carried out by us under the guidance of {GUIDE}, {GUIDE_DESIG}. "
        "The software, diagrams, test results and screenshots in this report were produced from the implementation developed by the project team.")
    R.p("To the best of our knowledge, this report has not been submitted, in part or in full, to any other university or institution for the award of "
        "any degree or diploma. Information taken from published or unpublished sources has been acknowledged in the text and listed in the References.",
        after=30)
    sign_row(R, TEAM[0][0], TEAM[1][0], f"Enrollment No. {TEAM[0][2]}", f"Enrollment No. {TEAM[1][2]}", after=24)
    sign_row(R, TEAM[2][0], "", f"Enrollment No. {TEAM[2][2]}", "", after=18)
    R.p("Date: ____________________", after=0)
    R.p("Place: Shirpur")


def acknowledgement(R):
    h = R.front("ACKNOWLEDGEMENT")
    h.runs[0].underline = True
    R.p("After the completion of this capstone project work, words are not enough to express our feelings about all those who helped us to reach our "
        "goal; above all, we are indebted to The Almighty for providing us this moment in life.")
    R.p(f"It is a great pleasure and a moment of immense satisfaction for us to express our profound gratitude to our project guide, **{GUIDE}**, "
        f"{GUIDE_DESIG}, MPSTME, Shirpur, whose constant encouragement enabled us to work enthusiastically. Their motivation, patience and expert "
        "guidance during the progress of the project benefited us beyond expression, and their depth of knowledge of Computer Science showed us that "
        "sound theory is the basis of efficient working software. We are also grateful to them, as Head of Department, for providing the infrastructure "
        "and support needed to complete the project successfully.")
    R.p(f"We extend our deepest gratitude to **{DEAN}**, Associate Dean, SVKM’s NMIMS, Shirpur Campus, for providing all the necessary facilities and "
        "an encouraging environment to bring out the best of our endeavours.")
    R.p("We sincerely thank all members of the staff of the Department of Computer Science, who imparted technical knowledge to us during the "
        "various stages of B.Tech. Computer Science, and the reviewers whose comments improved this report. We also thank our friends, who contributed "
        "directly or indirectly to this work, and the open-source communities whose tools made it possible.")
    R.p("", after=14)
    for name, _, _ in TEAM:
        R.p(name, align=WD_ALIGN_PARAGRAPH.RIGHT, after=0, ls=1.15)


def abstract(R):
    h = R.front("ABSTRACT")
    h.runs[0].underline = True
    R.p("People fill the same personal details into many web forms: admission portals, scholarship and examination applications, bank account forms "
        "and government services. The information already exists in documents they hold, such as identity cards, mark sheets and certificates, yet "
        "browser autofill can only repeat what was typed before. It cannot read documents, cannot show where a value came from, and does not warn when "
        "two documents disagree.")
    R.p("This report presents the Smart Form Auto-Filling System, a privacy-first assistant that builds a verified profile from a user’s own documents "
        "and uses it to suggest values for any web form. The system has a FastAPI backend, a React dashboard and a Chrome extension. Uploaded PDF, Word, "
        "text and image files are read (with optical character recognition for scans), classified into seventeen Indian document types, and turned into "
        "facts from a vocabulary of 220 concepts. Every fact keeps its source document, page and evidence text. Facts are stored in MongoDB and the "
        "document text is indexed in ChromaDB for semantic search. On a web form, a LangGraph pipeline maps each field to a concept, proposes a value "
        "from the profile or from cited document text, and decides whether the value is ready or must be reviewed. The user approves every value before "
        "the extension fills it, and the form is never submitted automatically. A question-answering page answers questions about the documents with "
        "cited sources.")
    R.p("On 22 October 2026 all 446 automated tests that were run passed (301 backend, 3 vector-database, 30 dashboard and 112 extension tests), and the "
        "7 browser end-to-end tests passed in the recorded run of 15 October 2026. On a synthetic set of 15 documents and 15 forms, 97.0% of expected "
        "fields were extracted correctly, all 28 test questions were answered correctly with sources, and 112 form fields were filled correctly, 9 "
        "incorrectly and 84 were held for the user’s review, with no form submitted.")
    p = R.p("")
    p.add_run("Keywords: ").bold = True
    p.add_run("form autofill, document information extraction, optical character recognition, retrieval-augmented generation, LangGraph, "
              "MongoDB, Chrome extension, privacy.")


ABBREVIATIONS = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"), ("CAPTCHA", "Test that tells computers and humans apart"),
    ("CI", "Continuous Integration"), ("CPU", "Central Processing Unit"), ("CSS", "Cascading Style Sheets"), ("DFD", "Data-Flow Diagram"),
    ("DOCX", "Microsoft Word document format"), ("ER", "Entity–Relationship"), ("FR", "Functional Requirement"),
    ("HTTP", "Hypertext Transfer Protocol"), ("ICLR / ICML", "Int. Conf. on Learning Representations / on Machine Learning"),
    ("JSON", "JavaScript Object Notation"), ("JWT", "JSON Web Token"), ("KYC", "Know Your Customer"), ("LLM", "Large Language Model"),
    ("MV3", "Manifest Version 3 (Chrome extension platform)"), ("OCR", "Optical Character Recognition"), ("ONNX", "Open Neural Network Exchange"),
    ("OTP", "One-Time Password"), ("OWASP", "Open Worldwide Application Security Project"), ("PDF", "Portable Document Format"),
    ("PIN", "Postal Index Number"), ("RAG", "Retrieval-Augmented Generation"), ("REST", "Representational State Transfer"),
    ("TC", "Test Case"), ("UI", "User Interface"), ("UML", "Unified Modelling Language"),
]


def abbreviations(R):
    R.front("LIST OF ABBREVIATIONS")
    from docx.enum.table import WD_TABLE_ALIGNMENT
    t = R.doc.add_table(rows=len(ABBREVIATIONS) + 1, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate([("Abbreviation", "Full form")] + ABBREVIATIONS):
        for j, v in enumerate(row):
            cell = t.rows[i].cells[j]
            cell.width = Inches(1.5 if j == 0 else 4.27)
            p = cell.paragraphs[0]
            p.style = "TableText"
            r = p.add_run(v)
            r.bold = i == 0
            r.font.size = Pt(11)
