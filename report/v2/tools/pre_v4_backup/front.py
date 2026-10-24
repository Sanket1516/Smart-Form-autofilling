# -*- coding: utf-8 -*-
"""Preliminary pages, laid out after the department's reference black book (Company_Brain_Final.pdf)."""
import os

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from docbuild2 import HERE, TEXT_W, style_grid

TITLE = "Agentic AI Framework for Context-Aware Semantic Understanding and Intelligent Assistance of Dynamic Web Forms"
SHORT = "Smart Form Auto-Filling System"
# (name, roll number, enrollment number) as recorded in the earlier submission
TEAM = [("Ashmit Patil", "246", "70552300048"), ("Sanket Mali", "265", "70552300067"), ("Pinak Salunke", "247", "70552300049")]
# Reviewer note: the guide's designation is Head of Department, not Professor.
GUIDE = "Dr. Nitin S. Choubey"
GUIDE_DESIG = "Head of Department, Dept. of Computer Science"
HOD = "Dr. Nitin S. Choubey"
DEAN = "Dr. Venkatadri Marriboyina"
DIRECTOR = "Dr. Sunita Patil"
SESSION = "2026–27"
LOGO = os.path.join(HERE, "..", "figures", "nmims_logo.png")
C = WD_ALIGN_PARAGRAPH.CENTER
L = WD_ALIGN_PARAGRAPH.LEFT
RT = WD_ALIGN_PARAGRAPH.RIGHT
DEPT = "DEPARTMENT OF COMPUTER SCIENCE"
SCHOOL = "Mukesh Patel School of Technology Management & Engineering"


def line(R, text, size=12, bold=False, after=6, italic=False, before=0, underline=False, align=C):
    p = R.p("", align=align, after=after, before=before, ls=1.15)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.underline = underline
    return p


def logo(R, width, after=6):
    p = R.p("", align=C, after=after, ls=1.0)
    p.add_run().add_picture(LOGO, width=Inches(width))


def plain_table(R, rows, widths, size=12, bold_cols=(), aligns=None, row_after=4):
    """Borderless layout table (names, enrollment numbers, signature lines)."""
    t = R.doc.add_table(rows=len(rows), cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = t.rows[i].cells[j]
            cell.width = Inches(widths[j])
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(row_after)
            p.paragraph_format.first_line_indent = Inches(0)
            p.alignment = aligns[j] if aligns else L
            for k, part in enumerate(str(v).split("\n")):
                if k:
                    p.add_run().add_break()
                r = p.add_run(part)
                r.font.size = Pt(size)
                r.bold = j in bold_cols
    return t


def dept_footer(R, before=24):
    line(R, DEPT, 14, True, after=2, before=before)
    line(R, SCHOOL, 12, True, after=0)


def cover(R):
    line(R, TITLE, 22, True, after=18, before=6)
    line(R, "Submitted In Partial Fulfillment of the Requirement", 12, True, after=0)
    line(R, "for the Award of Degree of Bachelor of Technology In", 12, True, after=0)
    line(R, "Computer Science", 12, True, after=8)
    line(R, "Submitted To", 11, True, after=4)
    logo(R, 1.55, after=4)
    line(R, "SVKM’s NMIMS", 11, True, after=2)
    line(R, SCHOOL, 11, True, after=2)
    line(R, "Shirpur Campus Maharashtra", 11, True, after=8)
    line(R, "Submitted By", 11, True, after=4)
    plain_table(R, [(n, e) for n, _, e in TEAM], [2.1, 1.6], size=13, bold_cols=(0, 1), row_after=2)
    line(R, "Under The Supervision of:", 11, True, after=2, before=10)
    line(R, GUIDE, 13, True, after=2)
    line(R, GUIDE_DESIG, 11, True, after=14)
    line(R, DEPT, 13, True, after=2)
    line(R, SCHOOL, 11, True, after=2)
    line(R, f"ACADEMIC SESSION {SESSION}", 11, True, after=0)


def certificate(R):
    R.front("CERTIFICATE", toc_title="Certificate")
    logo(R, 1.5, after=14)
    p = R.p("", ls=1.5)
    p.add_run(f"This is to certify that the project entitled {TITLE} has been done by :-")
    for name, roll, enr in TEAM:
        q = R.p("", align=C, after=0, ls=1.15)
        q.add_run(f"“{name} - {enr}” (Roll No.: {roll})").bold = True
    p = R.p("", ls=1.5, before=2)
    p.add_run("under my guidance and supervision & has been submitted").bold = True
    p.add_run(" in partial fulfillment of the degree of ")
    p.add_run("“Bachelor of Technology in Computer Science”").bold = True
    p.add_run(" of SVKM’S NMIMS (Deemed-to-be University), MPSTME Shirpur Campus (M.H.), India.")
    R.p("", after=18, ls=1.0)
    plain_table(R, [("____________\n" + GUIDE + "\nProject Mentor\n(Internal Guide)", "\n____________\nExaminer")],
                [TEXT_W / 2, TEXT_W / 2], size=12, bold_cols=(0, 1), aligns=[L, RT])
    R.p("Date:", bold=True, after=4, before=8, ls=1.15)
    R.p("Place: Shirpur", bold=True, after=0, ls=1.15)
    plain_table(R, [("", "____________\n" + HOD + "\nHead of the Department")], [TEXT_W / 2, TEXT_W / 2], size=12,
                bold_cols=(1,), aligns=[L, RT])
    dept_footer(R, before=14)


def approval(R):
    R.front("CERTIFICATE OF APPROVAL", toc_title="Certificate of Approval")
    names = ", ".join(f"{n} ({e})" for n, _, e in TEAM[:-1]) + f" and {TEAM[-1][0]} ({TEAM[-1][2]})"
    R.p(f"The project entitled {TITLE}, submitted by {names} of B.Tech. (Computer Science), Academic Session {SESSION}, MPSTME Shirpur "
        "Campus, SVKM’s NMIMS, has been examined and is approved in partial fulfillment of the requirements for the award of the degree of "
        "“**Bachelor of Technology in Computer Science**”.")
    R.p("The examiners record their approval below after the project presentation and viva-voce examination held on ____________________.",
        after=36)
    plain_table(R, [("____________________\nInternal Examiner\nName:", "____________________\nExternal Examiner\nName:")],
                [TEXT_W / 2, TEXT_W / 2], bold_cols=(0, 1), aligns=[L, RT])
    R.p("", after=24, ls=1.0)
    plain_table(R, [("____________________\n" + GUIDE + "\nProject Mentor", "____________________\n" + HOD + "\nHead of the Department")],
                [TEXT_W / 2, TEXT_W / 2], bold_cols=(0, 1), aligns=[L, RT])
    R.p("Date:", bold=True, after=4, before=14, ls=1.15)
    R.p("Place: Shirpur", bold=True, after=0, ls=1.15)
    dept_footer(R, before=60)


def declaration(R):
    R.front("DECLARATION", toc_title="Declaration")
    p = R.p("")
    p.add_run(f"We hereby declare that the project entitled {TITLE} submitted in partial fulfillment of the degree of ")
    p.add_run("“Bachelor of Technology in Computer Science”").bold = True
    p.add_run(" of Mukesh Patel School of Technology Management & Engineering Shirpur, SVKM’s Narsee Monjee Institute of Management Studies "
              "(NMIMS) Deemed-to-be University (M.H.), India, is a record of original work carried out by us ")
    p.add_run(f"under the guidance and supervision of {GUIDE},").bold = True
    p.add_run(" Head of Department, Department of Computer Science.")
    R.p("The matter embodied in this report has not been submitted, in part or fully, for the award of any other degree or diploma from this or "
        "any other university or institution. Wherever the work of others has been used, it has been duly acknowledged and cited.", after=40)
    plain_table(R, [(n, e, "________________________") for n, _, e in TEAM], [1.7, 1.5, 2.4], size=12, bold_cols=(0, 1), row_after=14)
    R.p("Date:", bold=True, after=4, before=4, ls=1.15)
    R.p("Place: Shirpur", bold=True, after=0, ls=1.15)
    dept_footer(R, before=90)


def acknowledgement(R):
    R.front("ACKNOWLEDGEMENT", toc_title="Acknowledgement")
    R.p("On completing this capstone project, we would like to thank everyone who helped us reach our goal. Above all, we are grateful to the "
        "Almighty for giving us this opportunity and guiding us throughout the project.")
    R.p(f"We are deeply grateful to **{GUIDE}**, Head of the Department of Computer Science, MPSTME, Shirpur, for their invaluable guidance, "
        "constant encouragement and support throughout the project. Their expertise and suggestions helped us put our theoretical knowledge "
        "into practice. As Head of the Department, they also provided the infrastructure and resources we needed.")
    R.p(f"We also extend our gratitude to **{DEAN}**, Associate Dean of MPSTME, Shirpur Campus, and **{DIRECTOR}**, Director of SVKM’s NMIMS, "
        "Shirpur Campus, for providing us with the necessary facilities, encouragement and a supportive learning environment.")
    R.p("We are thankful to the faculty members and staff of the Computer Science Department for the technical knowledge and support they gave "
        "us throughout our B.Tech. studies, and to the reviewers whose comments improved this report. We also thank our friends for their "
        "encouragement and their direct and indirect contributions to this project.")
    R.p("Finally, we sincerely thank everyone whose guidance, effort and support made the completion of this project possible.")
    R.p("", after=14)
    for name, _, _ in TEAM:
        R.p(name, align=RT, after=0, ls=1.15)


def abstract(R):
    R.front("ABSTRACT", toc_title="Abstract")
    R.p("People fill the same personal details into many web forms: admission portals, scholarship and examination applications, bank account forms "
        "and government services. The information already exists in documents they hold, such as identity cards, mark sheets and certificates, yet "
        "browser autofill can only repeat what was typed before. It cannot read documents, cannot show where a value came from, and does not warn when "
        "two documents disagree.")
    R.p("This report presents the Smart Form Auto-Filling System, a privacy-first assistant that builds a verified profile from a user’s own documents "
        "and uses it to suggest values for any web form. The system has a FastAPI backend, a React dashboard and a Chrome extension. Uploaded PDF, Word, "
        "text and image files are read (with optical character recognition for scans), classified into seventeen Indian document types, and turned into "
        "facts from a vocabulary of 220 concepts. Every fact keeps its source document, page and evidence text. Facts are stored in MongoDB and the "
        "document text is indexed in ChromaDB for semantic search. On a web form, a LangGraph pipeline maps each field to a concept, proposes a value "
        "from the profile or from cited document text, and decides whether the value is ready or must be reviewed. The user approves every value before "
        "the extension fills it, and the form is never submitted automatically. A question-answering page answers questions about the documents with "
        "cited sources.")
    R.p("On 22 October 2026 all 446 automated tests that were run passed (301 backend, 3 vector-database, 30 dashboard and 112 extension tests), and the "
        "7 browser end-to-end tests passed in the recorded run of 15 October 2026. On a synthetic set of 15 documents and 15 forms, 97.0% of expected "
        "fields were extracted correctly and all 28 test questions were answered correctly with sources. In form filling, 112 fields were filled "
        "correctly and 9 incorrectly, 84 were held for the user’s review, and no form was submitted.")
    R.p("*Keywords: form autofill; document information extraction; optical character recognition; retrieval-augmented generation; LangGraph; "
        "MongoDB; Chrome extension; privacy.*")


ABBREVIATIONS = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"), ("CAPTCHA", "Test that tells computers and humans apart"),
    ("CI", "Continuous Integration"), ("CPU", "Central Processing Unit"), ("CSS", "Cascading Style Sheets"), ("DFD", "Data-Flow Diagram"),
    ("DOCX", "Microsoft Word document format"), ("ER", "Entity–Relationship"), ("FR", "Functional Requirement"),
    ("HTTP", "Hypertext Transfer Protocol"), ("ICLR / ICML", "Int. Conf. on Learning Representations / on Machine Learning"),
    ("JSON", "JavaScript Object Notation"), ("JWT", "JSON Web Token"), ("KYC", "Know Your Customer"), ("LLM", "Large Language Model"),
    ("MV3", "Manifest Version 3 (Chrome extension platform)"), ("OCR", "Optical Character Recognition"), ("ONNX", "Open Neural Network Exchange"),
    ("OTP", "One-Time Password"), ("OWASP", "Open Worldwide Application Security Project"), ("PDF", "Portable Document Format"),
    ("PIN", "Postal Index Number"), ("RAG", "Retrieval-Augmented Generation"), ("REST", "Representational State Transfer"),
    ("TC", "Test Case"), ("UI", "User Interface"), ("UML", "Unified Modelling Language"),
]


def abbreviations(R):
    R.front("LIST OF ABBREVIATIONS", toc_title="List of Abbreviations", underline=False)
    t = R.doc.add_table(rows=len(ABBREVIATIONS) + 1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate([("Abbreviation", "Full form")] + ABBREVIATIONS):
        for j, v in enumerate(row):
            cell = t.rows[i].cells[j]
            cell.width = Inches(1.5 if j == 0 else 4.27)
            p = cell.paragraphs[0]
            p.style = "TableText"
            r = p.add_run(v)
            r.bold = i == 0
            r.font.size = Pt(11)
    style_grid(t, pad=10)
