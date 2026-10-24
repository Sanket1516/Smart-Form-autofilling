# -*- coding: utf-8 -*-
"""python-docx toolkit for the black book (report/v2).

Layout follows the department's reference black book (Company_Brain_Final.pdf) within the rules of
Report Writing Instructions.pdf:
  A4 with a double page border; Times New Roman; body 12 pt, line spacing 1.5, justified, first-line indent;
  margins L 1.5", R/T/B 1.0";
  "Chapter N" 24 pt bold under a blue rule, chapter name in italic capitals over a black rule;
  section headings 12 pt bold;
  tables with a dark header row and alternate shaded rows, caption above; framed figures, caption below;
  preliminary pages without a running header and with roman page numbers;
  chapters with a grey running header (project name left, academic session right) and decimal page numbers;
  references and appendices continue the decimal numbering without the running header;
  table of contents, list of figures and list of tables are tables whose page numbers are PAGEREF fields.

Figures and tables are numbered per chapter in the order they are first referenced or inserted,
and a figure is kept on the same page as the paragraph that introduces it.
"""
import os
import re

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
FONT = "Times New Roman"
CODE_FONT = "Consolas"
PAGE_W, PAGE_H = 8.27, 11.69
ML, MR, MT, MB = 1.5, 1.0, 1.0, 1.0
TEXT_W = round(PAGE_W - ML - MR, 2)  # 5.77
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x76, 0x76, 0x76)
HEADER_LEFT = "SMART FORM AUTO-FILLING SYSTEM"
HEADER_RIGHT = "ACADEMIC SESSION 2026–27"  # same form as the cover page
INDENT = Inches(0.5)
CODE_COLS = 82  # characters of 9 pt Consolas that fit on one line of the text block
HEAD_FILL, ZEBRA_FILL, GRID = "3F3F3F", "F2F2F2", "BFBFBF"


def set_font(obj, name=FONT, size=None, bold=None, italic=None, color=BLACK):
    f = obj.font
    f.name = name
    rpr = obj.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def add_field(paragraph, instr, placeholder="1"):
    def fc(t):
        r = paragraph.add_run()
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), t)
        r._r.append(e)
        return r
    fc("begin")
    r = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)
    fc("separate")
    paragraph.add_run(placeholder)
    fc("end")


def para_border(paragraph, side, size=6, color="000000", space=1):
    ppr = paragraph._p.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        ppr.append(bdr)
    sides = [side] if side != "box" else ["top", "left", "bottom", "right"]
    for sd in sides:
        b = OxmlElement(f"w:{sd}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(size))
        b.set(qn("w:space"), str(space))
        b.set(qn("w:color"), color)
        bdr.append(b)


def bottom_border(paragraph, size=6):
    para_border(paragraph, "bottom", size)


def style_grid(tbl, header=True, zebra=True, pad=40):
    """Light grey grid, dark header row with white bold text, alternate shaded body rows, padded cells."""
    tblpr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), GRID)
        borders.append(b)
    tblpr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", pad), ("bottom", pad), ("left", 80), ("right", 80)):
        m = OxmlElement(f"w:{side}")
        m.set(qn("w:w"), str(v))
        m.set(qn("w:type"), "dxa")
        mar.append(m)
    tblpr.append(mar)
    for ri, row in enumerate(tbl.rows):
        for cell in row.cells:
            if header and ri == 0:
                shade(cell, HEAD_FILL)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.color.rgb = WHITE
                        r.bold = True
            elif zebra and ri % 2 == 0:
                shade(cell, ZEBRA_FILL)


class Report:
    def __init__(self, figures, tables):
        self.doc = Document()
        self.figures = figures      # key -> dict(file, caption, chapter)
        self.tables = tables        # key -> dict(caption, chapter)
        self.num = {}               # ("f"|"t", key) -> "4.2"
        self.counter = {}
        self.inserted = []          # (kind, key, number) in insertion order
        self.referenced = set()
        self.indent_body = False    # first-line indent for body paragraphs (chapters and appendices)
        self.toc_entries = []       # (level, number, title, bookmark); level 0 = preliminary page or back matter
        self.lof, self.lot = [], []  # (number, caption, bookmark)
        self.placeholders = {}      # "toc" | "lof" | "lot" -> paragraph to replace with the generated table
        self._bm = 0
        self._setup()

    # ------------------------------------------------------------------ setup
    def _setup(self):
        d = self.doc
        st = d.styles
        n = st["Normal"]
        set_font(n, size=12)
        pf = n.paragraph_format
        pf.line_spacing = 1.5
        pf.space_after = Pt(6)
        pf.space_before = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.widow_control = True

        h1 = st["Heading 1"]
        set_font(h1, size=24, bold=True)
        h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h1.paragraph_format.space_before = Pt(6)
        h1.paragraph_format.space_after = Pt(18)
        h1.paragraph_format.line_spacing = 1.15
        h1.paragraph_format.keep_with_next = True
        h1.paragraph_format.page_break_before = True
        for name, before, after in (("Heading 2", 10, 4), ("Heading 3", 8, 2)):
            h = st[name]
            set_font(h, size=12, bold=True, italic=False)
            h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.line_spacing = 1.5

        def pstyle(name, size=12, bold=False, italic=False, align=None, ls=1.5, before=0, after=6, keep_next=False, font=FONT):
            s_ = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            s_.base_style = st["Normal"]
            set_font(s_, name=font, size=size, bold=bold, italic=italic)
            s_.paragraph_format.line_spacing = ls
            s_.paragraph_format.space_before = Pt(before)
            s_.paragraph_format.space_after = Pt(after)
            s_.paragraph_format.first_line_indent = Inches(0)
            if align is not None:
                s_.paragraph_format.alignment = align
            s_.paragraph_format.keep_with_next = keep_next
            return s_

        fh = pstyle("FrontHeading", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, ls=1.15, after=18, keep_next=True)
        fh.paragraph_format.page_break_before = True
        pstyle("ChapterTitle", size=18, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, ls=1.15, before=0, after=14, keep_next=True)
        pstyle("FigureCaption", size=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, ls=1.0, before=4, after=10)
        pstyle("TableCaption", size=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, ls=1.0, before=8, after=4, keep_next=True)
        pstyle("TableText", size=10.5, ls=1.0, after=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        pstyle("CodeBlock", size=9, bold=True, ls=1.0, after=0, align=WD_ALIGN_PARAGRAPH.LEFT, font=CODE_FONT)
        lb = st["List Bullet"]
        set_font(lb, size=12)
        lb.paragraph_format.line_spacing = 1.5
        lb.paragraph_format.space_after = Pt(2)
        lb.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        self._page(d.sections[0])

    def _page(self, s):
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Inches(PAGE_W), Inches(PAGE_H)
        s.left_margin, s.right_margin = Inches(ML), Inches(MR)
        s.top_margin, s.bottom_margin = Inches(MT), Inches(MB)
        s.header_distance = Inches(0.5)
        s.footer_distance = Inches(0.5)

    def _sect_props(self, s, fmt, start):
        """Page border and page numbering, inserted in the order the schema requires (after w:pgMar)."""
        sect = s._sectPr
        for tag in ("w:pgBorders", "w:pgNumType"):
            for e in sect.findall(qn(tag)):
                sect.remove(e)
        pgmar = sect.find(qn("w:pgMar"))
        borders = OxmlElement("w:pgBorders")
        borders.set(qn("w:offsetFrom"), "page")
        for side in ("top", "left", "bottom", "right"):
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "double")
            b.set(qn("w:sz"), "6")
            b.set(qn("w:space"), "24")
            b.set(qn("w:color"), "000000")
            borders.append(b)
        pg = OxmlElement("w:pgNumType")
        pg.set(qn("w:fmt"), fmt)
        if start is not None:
            pg.set(qn("w:start"), str(start))
        pgmar.addnext(borders)
        borders.addnext(pg)

    def _clear(self, part):
        for p in part.paragraphs:
            for r in list(p.runs):
                r._r.getparent().remove(r._r)

    def header_footer(self, s, mode="body"):
        """mode: "none" (cover), "plain" (page number only) or "body" (running header and page number)."""
        s.header.is_linked_to_previous = False
        s.footer.is_linked_to_previous = False
        self._clear(s.header)
        self._clear(s.footer)
        if mode == "none":
            return
        if mode == "body":
            hp = s.header.paragraphs[0]
            hp.paragraph_format.line_spacing = 1.0
            hp.paragraph_format.space_after = Pt(0)
            hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            ts = hp.paragraph_format.tab_stops
            for pos in (3.25, 6.5):  # clear the centre and right tabs inherited from Word's Header style
                ts.add_tab_stop(Inches(pos), WD_TAB_ALIGNMENT.CLEAR)
            ts.add_tab_stop(Inches(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
            r = hp.add_run(f"{HEADER_LEFT}\t{HEADER_RIGHT}")
            set_font(r, size=10, color=GREY)
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp.paragraph_format.line_spacing = 1.0
        fp.paragraph_format.space_after = Pt(0)
        add_field(fp, "PAGE", "1")
        for r in fp.runs:
            set_font(r, size=12)

    def first_section(self):
        s = self.doc.sections[0]
        self._sect_props(s, "lowerRoman", 1)
        self.header_footer(s, "none")

    def new_section(self, fmt, start, mode="body"):
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._page(s)
        self._sect_props(s, fmt, start)
        self.header_footer(s, mode)
        return s

    # ------------------------------------------------------------------ bookmarks
    def mark(self, paragraph, name=None):
        self._bm += 1
        name = name or f"_Ref_sf{self._bm:03d}"
        p = paragraph._p
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(self._bm))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(self._bm))
        p.insert(1 if p.pPr is not None else 0, start)
        p.append(end)
        return name

    # ------------------------------------------------------------------ text
    def p(self, text="", style=None, align=None, bold=False, italic=False, keep_next=False, size=None, after=None, before=None, ls=None,
          indent=None):
        para = self.doc.add_paragraph(style=style)
        if text:
            self.runs(para, text, bold, italic, size)
        if align is not None:
            para.alignment = align
        pf = para.paragraph_format
        if indent is None:
            indent = self.indent_body and style is None and align is None
        if indent:
            pf.first_line_indent = INDENT
        if keep_next:
            pf.keep_with_next = True
        if after is not None:
            pf.space_after = Pt(after)
        if before is not None:
            pf.space_before = Pt(before)
        if ls is not None:
            pf.line_spacing = ls
        return para

    def runs(self, para, text, bold=False, italic=False, size=None):
        """Inline markup: **bold**, *italic*, `name` (identifier, set in italic Times New Roman)."""
        for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
            if not part:
                continue
            if part.startswith("**"):
                r = para.add_run(part[2:-2])
                r.bold = True
                r.italic = italic or None
            elif part.startswith("`"):
                r = para.add_run(part[1:-1])
                r.italic = True
                r.bold = bold or None
            elif part.startswith("*") and len(part) > 2:
                r = para.add_run(part[1:-1])
                r.italic = True
                r.bold = bold or None
            else:
                r = para.add_run(part)
                r.bold = bold or None
                r.italic = italic or None
            if size:
                r.font.size = Pt(size)

    def chapter(self, num, title):
        self.chapter_no = num
        head = self.doc.add_paragraph(style="Heading 1")
        head.paragraph_format.space_before = Pt(0)
        head.paragraph_format.space_after = Pt(0)
        para_border(head, "top", size=6, color="5B9BD5", space=10)
        head.add_run(f"Chapter {num}")
        name = self.mark(head, "BodyStart" if num == 1 else None)
        sub = self.doc.add_paragraph(title.upper(), style="ChapterTitle")
        para_border(sub, "bottom", size=12, space=2)
        self.toc_entries.append((1, str(num), title.upper(), name))
        return head

    def h1_plain(self, title, toc=True):
        para = self.doc.add_paragraph(title, style="Heading 1")
        if toc:
            self.toc_entries.append((0, "", title, self.mark(para)))
        return para

    def front(self, title, toc_title=None, underline=True):
        para = self.doc.add_paragraph(title, style="FrontHeading")
        if underline:
            para.runs[0].underline = True
        if toc_title:
            self.toc_entries.append((0, "", toc_title, self.mark(para)))
        return para

    def _heading(self, text, style, level):
        para = self.doc.add_paragraph(text, style=style)
        m = re.match(r"(\d+(?:\.\d+)+)\s+(.*)", text)
        if m:
            self.toc_entries.append((level, m.group(1), m.group(2), self.mark(para)))
        return para

    def h2(self, text):
        return self._heading(text, "Heading 2", 2)

    def h3(self, text):
        return self._heading(text, "Heading 3", 3)

    def bullets(self, items, keep_last_next=False):
        out = []
        for it in items:
            para = self.doc.add_paragraph(style="List Bullet")
            if "::" in it:
                a, b = it.split("::", 1)
                r = para.add_run(a.strip() + ": ")
                r.bold = True
                self.runs(para, b.strip())
            else:
                self.runs(para, it)
            out.append(para)
        if keep_last_next and out:
            out[-1].paragraph_format.keep_with_next = True
        return out

    def numbered(self, items, bold_lead=True):
        for i, it in enumerate(items, 1):
            para = self.doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.35)
            para.paragraph_format.first_line_indent = Inches(-0.3)
            para.paragraph_format.space_after = Pt(2)
            # a tab (not spaces) after the number, so justification cannot widen it and every item starts at the same position
            para.paragraph_format.tab_stops.add_tab_stop(Inches(0.35))
            if bold_lead and "::" in it:
                a, b = it.split("::", 1)
                r = para.add_run(f"{i}.\t{a.strip()}: ")
                r.bold = True
                self.runs(para, b.strip())
            else:
                self.runs(para, f"{i}.\t" + it)

    def summary(self, text):
        self.h2("Chapter Summary")
        return self.p(text)

    # ------------------------------------------------------------------ numbering
    def _number(self, kind, key):
        if (kind, key) not in self.num:
            meta = (self.figures if kind == "f" else self.tables)[key]
            ch = meta["chapter"]
            self.counter[(kind, ch)] = self.counter.get((kind, ch), 0) + 1
            self.num[(kind, key)] = f"{ch}.{self.counter[(kind, ch)]}"
        return self.num[(kind, key)]

    def f(self, key):
        self.referenced.add(("f", key))
        return "Figure " + self._number("f", key)

    def t(self, key):
        self.referenced.add(("t", key))
        return "Table " + self._number("t", key)

    # ------------------------------------------------------------------ figures
    def fig(self, key, width=None, max_h=None, intro=None):
        """Insert a framed figure. ``intro`` is the explanatory paragraph placed directly above it and kept on the same page."""
        meta = self.figures[key]
        number = self._number("f", key)
        self.inserted.append(("f", key, number))
        if intro:
            self.p(intro, keep_next=True)
        path = os.path.join(HERE, "figures", meta["file"])
        w_px, h_px = Image.open(path).size
        aspect = w_px / h_px
        w = min(width or TEXT_W, TEXT_W)
        if max_h:
            w = min(w, max_h * aspect)
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = para.paragraph_format
        pf.keep_with_next = True
        pf.line_spacing = 1.0
        pf.space_before = Pt(4)
        pf.space_after = Pt(0)
        run = para.add_run()
        run.add_picture(path, width=Inches(w))
        sppr = run._r.find(".//" + qn("pic:spPr"))
        ln = OxmlElement("a:ln")
        ln.set("w", "9525")
        fill = OxmlElement("a:solidFill")
        clr = OxmlElement("a:srgbClr")
        clr.set("val", "000000")
        fill.append(clr)
        ln.append(fill)
        sppr.append(ln)
        cp = self.doc.add_paragraph(f"Figure {number} {meta['caption']}", style="FigureCaption")
        self.lof.append((number, meta["caption"], self.mark(cp)))
        return cp

    # ------------------------------------------------------------------ tables
    def table(self, key, rows, widths, font=10.5, align=None, keep=True, bold_first_col=False, intro=None):
        meta = self.tables[key]
        number = self._number("t", key)
        self.inserted.append(("t", key, number))
        if intro:
            self.p(intro, keep_next=True)
        total = sum(widths)
        if total > TEXT_W + 0.01:
            widths = [w * TEXT_W / total for w in widths]
        cap = self.doc.add_paragraph(f"Table {number} {meta['caption']}", style="TableCaption")
        self.lot.append((number, meta["caption"], self.mark(cap)))
        tbl = self.doc.add_table(rows=len(rows), cols=len(widths))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tbl._tbl.tblPr.append(lay)
        for ri, row in enumerate(rows):
            tr = tbl.rows[ri]
            trpr = tr._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if ri == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            last = ri == len(rows) - 1
            for ci, val in enumerate(row):
                cell = tr.cells[ci]
                cell.width = Inches(widths[ci])
                lines = str(val).split("\n")
                for li, line in enumerate(lines):
                    para = cell.paragraphs[0] if li == 0 else cell.add_paragraph()
                    para.style = "TableText"
                    if keep and not last:
                        para.paragraph_format.keep_with_next = True
                    if align and align[ci] == "c":
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    is_bold = ri == 0 or (bold_first_col and ci == 0)
                    self.runs(para, line, bold=is_bold, size=font)
        style_grid(tbl)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        spacer.paragraph_format.line_spacing = 0.5
        return tbl

    def code(self, text, caption=None):
        if caption:
            self.p(caption, bold=True, keep_next=True, after=2)
        lines = text.rstrip("\n").split("\n")
        for i, ln in enumerate(lines):
            para = self.doc.add_paragraph(style="CodeBlock")
            para.paragraph_format.left_indent = Inches(0.05)
            if len(ln) > CODE_COLS:
                # A line wider than the text block wraps after a "|" or "," (zero-width break, the code text is unchanged)
                # and its continuation is indented, instead of breaking inside a word.
                ln = re.sub(r"([|,])(?=\S)", "\\1\u200b", ln)
                para.paragraph_format.left_indent = Inches(0.45)
                para.paragraph_format.first_line_indent = Inches(-0.4)
            if i < min(3, len(lines) - 1):  # keep the first lines with the explanation; long listings may continue on the next page
                para.paragraph_format.keep_with_next = True
            run = para.add_run(ln if ln else " ")
            set_font(run, name=CODE_FONT, size=9, bold=True)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------------ contents and lists
    def list_placeholder(self, kind):
        """Mark where the table of contents ("toc") or the list of figures/tables ("lof"/"lot") goes; filled in by save()."""
        self.placeholders[kind] = self.doc.add_paragraph()

    def _list_table(self, header, rows, widths, merge_first=False):
        """rows: (cells, bold, bookmark). The last column holds a PAGEREF field that Word fills in."""
        ncol = len(widths)
        tbl = self.doc.add_table(rows=len(rows) + 1, cols=ncol)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tbl._tbl.tblPr.append(lay)
        for ri in range(len(rows) + 1):
            tr = tbl.rows[ri]
            trpr = tr._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if ri == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            for ci in range(ncol):
                tr.cells[ci].width = Inches(widths[ci])
        hdr = tbl.rows[0].cells
        if merge_first:
            hdr[0].merge(hdr[1])
        for ci, text in enumerate(header):
            cell = hdr[0] if (merge_first and ci == 0) else hdr[ci + (1 if merge_first and ci else 0)]
            para = cell.paragraphs[0]
            para.style = "TableText"
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = para.add_run(text)
            set_font(r, size=11, bold=True)
        for ri, (cells, bold, bm) in enumerate(rows, 1):
            for ci in range(ncol):
                para = tbl.rows[ri].cells[ci].paragraphs[0]
                para.style = "TableText"
                if ci == ncol - 1:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_field(para, f"PAGEREF {bm} \\h", "0")
                    for r in para.runs:
                        set_font(r, size=11, bold=bold)
                    continue
                text = cells[ci]
                if text:
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT if (merge_first and ci < 2) else (
                        WD_ALIGN_PARAGRAPH.CENTER if ci == 0 else WD_ALIGN_PARAGRAPH.LEFT)
                    r = para.add_run(text)
                    set_font(r, size=11, bold=bold)
        style_grid(tbl, header=False, zebra=False, pad=15)
        return tbl

    def _build_lists(self):
        made = {}
        if "toc" in self.placeholders:
            rows = []
            for level, number, title, bm in self.toc_entries:
                if level == 0:
                    rows.append((["", "", title], number == "" and not title.startswith(("Certificate", "Declaration", "Acknowledgement",
                                                                                        "Abstract", "List of")), bm))
                elif level == 1:
                    rows.append((["", number, title], True, bm))
                else:
                    rows.append((["", number, title], False, bm))
            made["toc"] = self._list_table(["Sr. No.", "Chapter No.", "Page"], rows, [0.35, 0.6, 4.02, 0.8], merge_first=True)
        if "lof" in self.placeholders:
            made["lof"] = self._list_table(["Figure No.", "Title", "Page"], [([n, c], False, bm) for n, c, bm in self.lof],
                                           [0.95, 4.02, 0.8])
        if "lot" in self.placeholders:
            made["lot"] = self._list_table(["Table No.", "Title", "Page"], [([n, c], False, bm) for n, c, bm in self.lot],
                                           [0.95, 4.02, 0.8])
        for kind, tbl in made.items():
            ph = self.placeholders[kind]._p
            ph.addprevious(tbl._tbl)
            ph.getparent().remove(ph)

    # ------------------------------------------------------------------ checks
    def check(self):
        problems = []
        seen = {}
        for kind, key, number in self.inserted:
            if (kind, key) in seen:
                problems.append(f"{kind} {key} inserted twice")
            seen[(kind, key)] = number
        for kind in ("f", "t"):
            order = [n for k, _, n in self.inserted if k == kind]
            by_ch = {}
            for n in order:
                ch, i = n.rsplit(".", 1)
                by_ch.setdefault(ch, []).append(int(i))
            for ch, nums in by_ch.items():
                if nums != list(range(1, len(nums) + 1)):
                    problems.append(f"{kind} numbering out of order in chapter {ch}: {nums}")
        for kind, key in self.num:
            if (kind, key) not in seen:
                problems.append(f"{kind} {key} referenced but never inserted")
        for kind, key, _ in self.inserted:
            if (kind, key) not in self.referenced:
                problems.append(f"{kind} {key} inserted but never referenced in the text")
        return problems

    def save(self, path):
        self._build_lists()
        self.doc.save(path)
